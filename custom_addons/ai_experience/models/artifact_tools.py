import base64
import csv
import io
import json

from odoo import models
from odoo.addons.llm_tool.decorators import llm_tool

from .artifact_policy import (
    validate_artifact_output,
    validate_artifact_payload,
)


class AiArtifactTools(models.AbstractModel):
    _inherit = "llm.tool"

    @llm_tool(destructive_hint=True)
    def generate_artifact(self, artifact_type: str, title: str, rows_json: str) -> dict:
        """Create a user-owned enterprise artifact from structured rows.

        Supported types: csv, xlsx, pdf, docx, pptx, svg, json. The tool is
        intentionally separate from generic ORM access and is protected by the
        central risk registry. The resulting attachment belongs to the current
        user and is returned as a semantic artifact reference.
        """
        self.env["ai.gateway.execution.gate"].authorize("generate_artifact", context_label="artifact generation")
        try:
            payload = validate_artifact_payload(json.loads(rows_json or "{}"))
        except (TypeError, ValueError, json.JSONDecodeError):
            return {"error": "invalid or oversized artifact payload"}
        rows = payload.get("rows") or []
        columns = payload.get("columns") or (list(rows[0].keys()) if rows and isinstance(rows[0], dict) else [])
        kind = (artifact_type or "csv").lower()
        title = "".join(c if c.isalnum() or c in "-_ ." else "_" for c in (title or "artifact")).strip() or "artifact"

        if kind == "csv":
            out = io.StringIO(); w = csv.writer(out); w.writerow(columns)
            for row in rows: w.writerow([row.get(c, "") for c in columns])
            data, mimetype, filename = out.getvalue().encode("utf-8-sig"), "text/csv", f"{title}.csv"
        elif kind == "json":
            data, mimetype, filename = json.dumps(payload, ensure_ascii=False, indent=2).encode(), "application/json", f"{title}.json"
        elif kind == "svg":
            values = payload.get("values") or [float(r.get("value", 0)) for r in rows]
            labels = payload.get("labels") or [str(r.get("label", i+1)) for i, r in enumerate(rows)]
            max_v = max(values or [1]) or 1; width, height = 900, 480; base_y = height - 70; bw = max(10, int((width-80)/max(1,len(values))-8))
            svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}"><rect width="100%" height="100%" fill="#0f1115"/><text x="40" y="35" fill="#e8eaed" font-size="22">{title}</text>']
            for i,v in enumerate(values):
                x=40+i*(bw+8); h=int((height-150)*float(v)/max_v); y=base_y-h; label=str(labels[i])[:18] if i<len(labels) else str(i+1)
                svg += [f'<rect x="{x}" y="{y}" width="{bw}" height="{h}" rx="5" fill="#4f8cff"/>', f'<text x="{x+bw/2}" y="{base_y+22}" text-anchor="middle" fill="#9aa1ac" font-size="12">{label}</text>']
            svg.append('</svg>'); data=''.join(svg).encode(); mimetype='image/svg+xml'; filename=f'{title}.svg'
        elif kind == "xlsx":
            from openpyxl import Workbook
            bio=io.BytesIO(); wb=Workbook(); ws=wb.active; ws.title=(title[:31] or "Report")
            ws.append(columns)
            for row in rows: ws.append([row.get(c, "") for c in columns])
            wb.save(bio); data=bio.getvalue(); mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"; filename=f"{title}.xlsx"
        elif kind == "pdf":
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfgen import canvas
            bio=io.BytesIO(); c=canvas.Canvas(bio,pagesize=A4); width,height=A4; c.setFont("Helvetica-Bold",16); c.drawString(40,height-50,title[:90]); y=height-80; c.setFont("Helvetica",9)
            for row in rows[:45]: c.drawString(40,y," | ".join(str(row.get(k,"")) for k in columns)[:120]); y-=14
            c.save(); data=bio.getvalue(); mimetype="application/pdf"; filename=f"{title}.pdf"
        elif kind == "docx":
            from docx import Document
            doc=Document(); doc.add_heading(title,level=1); table=doc.add_table(rows=1,cols=max(1,len(columns)))
            for i,c in enumerate(columns): table.rows[0].cells[i].text=str(c)
            for row in rows:
                cells=table.add_row().cells
                for i,c in enumerate(columns): cells[i].text=str(row.get(c,""))
            bio=io.BytesIO(); doc.save(bio); data=bio.getvalue(); mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"; filename=f"{title}.docx"
        elif kind == "pptx":
            from pptx import Presentation
            prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text=title[:80]
            box=slide.shapes.add_textbox(500000,1300000,8500000,4500000); tf=box.text_frame
            for i,row in enumerate(rows[:25]):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=" | ".join(str(row.get(k,"")) for k in columns)
            bio=io.BytesIO(); prs.save(bio); data=bio.getvalue(); mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation"; filename=f"{title}.pptx"
        else:
            return {"error": f"unsupported artifact type: {kind}"}

        try:
            validate_artifact_output(data)
        except ValueError:
            return {"error": "generated artifact exceeds the output limit"}
        attachment = self.env["ir.attachment"].sudo().create({
            "name": filename, "datas": base64.b64encode(data).decode(), "mimetype": mimetype,
            "res_model": "res.users", "res_id": self.env.user.id, "public": False,
        })
        return {"status":"created", "filename":filename, "mimetype":mimetype, "attachment_id":attachment.id, "download_path":f"/api/artifacts/{attachment.id}", "size":len(data)}
