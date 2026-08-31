import base64
import csv
import io
import json
import logging
import os
import tempfile
from datetime import datetime
from urllib.parse import quote as _quote_filename

from odoo import http
from odoo.http import request

from odoo.addons.ai_semantic_api.controllers.semantic_api import _require_auth, _json_response, _require_privileged
from odoo.addons.ai_gateway.controllers.gateway import _audit, _cors_preflight_response
from odoo.addons.ai_gateway.controllers.file_policy import validate_upload, MAX_UPLOAD_BYTES
from odoo.addons.ai_gateway.controllers.output_firewall import scrub_public_text

_logger = logging.getLogger(__name__)


def _model(env, name):
    return env[name] if name in env else None



ROLE_AGENT_MAP = {
    "ai_business_tools.role_executive": "Executive Agent",
    "ai_business_tools.role_hr_manager": "HR Agent",
    "ai_business_tools.role_hr_staff": "HR Agent",
    "ai_business_tools.role_finance_manager": "Finance Agent",
    "ai_business_tools.role_finance_staff": "Finance Agent",
    "ai_business_tools.role_warehouse_manager": "Warehouse Agent",
    "ai_business_tools.role_warehouse_staff": "Warehouse Agent",
    "ai_business_tools.role_project_manager": "Project Agent",
    "ai_business_tools.role_manager": "Manager Agent",
    "ai_business_tools.role_employee": "Employee Agent",
}

def _assigned_agent(env, user):
    for xmlid, name in ROLE_AGENT_MAP.items():
        group = env.ref(xmlid, raise_if_not_found=False)
        if group and group in user.groups_id:
            return name
    return "Company Assistant"


def _safe_name(value, fallback="artifact"):
    value = "".join(c if c.isalnum() or c in "-_ ." else "_" for c in (value or ""))
    return value.strip() or fallback


def _artifact_bytes(kind, title, payload):
    """Generate common enterprise artifacts without coupling the UI to a library.
    CSV/JSON/SVG are dependency-free. XLSX/PDF/DOCX/PPTX use optional libraries
    already common in production Python environments; a clear 501 is returned if
    a requested optional generator is not installed."""
    kind = kind.lower()
    title = _safe_name(title)
    rows = payload.get("rows") or []
    columns = payload.get("columns") or (list(rows[0].keys()) if rows and isinstance(rows[0], dict) else [])

    if kind == "csv":
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(columns)
        for row in rows:
            writer.writerow([row.get(c, "") if isinstance(row, dict) else "" for c in columns])
        return f"{title}.csv", "text/csv", out.getvalue().encode("utf-8-sig")

    if kind == "json":
        return f"{title}.json", "application/json", json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")

    if kind == "svg":
        values = payload.get("values") or [float(r.get("value", 0)) for r in rows if isinstance(r, dict)]
        labels = payload.get("labels") or [str(r.get("label", i + 1)) for i, r in enumerate(rows) if isinstance(r, dict)]
        width, height = 900, 480
        max_v = max(values or [1]) or 1
        bar_w = max(10, int((width - 80) / max(1, len(values)) - 8))
        parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
                 '<rect width="100%" height="100%" fill="#0f1115"/>', f'<text x="40" y="35" fill="#e8eaed" font-size="22">{title}</text>']
        base_y = height - 70
        for i, value in enumerate(values):
            x = 40 + i * (bar_w + 8)
            h = int((height - 150) * float(value) / max_v)
            y = base_y - h
            label = str(labels[i])[:18] if i < len(labels) else str(i + 1)
            parts.append(f'<rect x="{x}" y="{y}" width="{bar_w}" height="{h}" rx="5" fill="#4f8cff"/>')
            parts.append(f'<text x="{x + bar_w/2}" y="{base_y + 22}" text-anchor="middle" fill="#9aa1ac" font-size="12">{label}</text>')
        parts.append('</svg>')
        return f"{title}.svg", "image/svg+xml", "".join(parts).encode("utf-8")

    if kind == "xlsx":
        try:
            import xlsxwriter
        except ImportError:
            raise RuntimeError("XLSX generator is not installed")
        bio = io.BytesIO()
        wb = xlsxwriter.Workbook(bio, {"in_memory": True})
        ws = wb.add_worksheet(title[:31] or "Report")
        for c, name in enumerate(columns):
            ws.write(0, c, name)
        for r, row in enumerate(rows, start=1):
            for c, name in enumerate(columns):
                ws.write(r, c, row.get(name, "") if isinstance(row, dict) else "")
        ws.freeze_panes(1, 0)
        wb.close()
        return f"{title}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", bio.getvalue()

    if kind == "pdf":
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfgen import canvas
        except ImportError:
            raise RuntimeError("PDF generator is not installed")
        bio = io.BytesIO()
        c = canvas.Canvas(bio, pagesize=A4)
        width, height = A4
        c.setFont("Helvetica-Bold", 16)
        c.drawString(40, height - 50, title[:90])
        y = height - 80
        c.setFont("Helvetica", 9)
        for row in rows[:45]:
            line = " | ".join(str(row.get(k, "")) for k in columns) if isinstance(row, dict) else str(row)
            c.drawString(40, y, line[:120])
            y -= 14
            if y < 45:
                c.showPage(); y = height - 50; c.setFont("Helvetica", 9)
        c.save()
        return f"{title}.pdf", "application/pdf", bio.getvalue()

    if kind == "docx":
        try:
            from docx import Document
        except ImportError:
            raise RuntimeError("DOCX generator is not installed")
        doc = Document(); doc.add_heading(title, level=1)
        if columns:
            table = doc.add_table(rows=1, cols=len(columns))
            for i, c in enumerate(columns): table.rows[0].cells[i].text = str(c)
            for row in rows:
                cells = table.add_row().cells
                for i, c in enumerate(columns): cells[i].text = str(row.get(c, ""))
        bio = io.BytesIO(); doc.save(bio)
        return f"{title}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", bio.getvalue()

    if kind == "pptx":
        try:
            from pptx import Presentation
        except ImportError:
            raise RuntimeError("PPTX generator is not installed")
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title[:80]
        box = slide.shapes.add_textbox(500000, 1300000, 8500000, 4500000)
        tf = box.text_frame
        for i, row in enumerate(rows[:25]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = " | ".join(str(row.get(k, "")) for k in columns) if isinstance(row, dict) else str(row)
        bio = io.BytesIO(); prs.save(bio)
        return f"{title}.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", bio.getvalue()

    raise ValueError(f"unsupported artifact type: {kind}")


class AiExperienceApi(http.Controller):
    @http.route("/api/workspace", type="http", auth="none", csrf=False, methods=["GET", "OPTIONS"])
    def workspace(self, **kwargs):
        if request.httprequest.method == "OPTIONS": return _cors_preflight_response()
        env, err = _require_auth()
        if err: return err
        user = env.user
        employee_model = _model(env, "hr.employee")
        employee = employee_model.search([("user_id", "=", user.id)], limit=1) if employee_model is not None else None
        capabilities = env["ai.control.capability.resolver"].effective_capabilities(user=user) if "ai.control.capability.resolver" in env else []
        return _json_response({
            "user": {"id": user.id, "name": user.name, "login": user.login},
            "company": {"id": env.company.id, "name": env.company.name},
            "department": {"id": employee.department_id.id, "name": employee.department_id.name} if employee and employee.department_id else None,
            "agent": _assigned_agent(env, user),
            "capabilities": capabilities,
        })

    @http.route("/api/departments", type="http", auth="none", csrf=False, methods=["GET", "OPTIONS"])
    def departments(self, **kwargs):
        if request.httprequest.method == "OPTIONS": return _cors_preflight_response()
        env, err = _require_auth()
        if err: return err
        dept_model = _model(env, "hr.department")
        if dept_model is None: return _json_response({"departments": []})
        employee_model = _model(env, "hr.employee")
        employee = employee_model.search([("user_id", "=", env.user.id)], limit=1) if employee_model is not None else None
        privileged = env.user.has_group("base.group_system")
        depts = dept_model.search([]) if privileged else (employee.department_id | employee.child_ids.mapped("department_id"))
        return _json_response({"departments": [{"id": d.id, "name": d.name, "manager_id": d.manager_id.id if d.manager_id else None, "member_count": len(d.member_ids)} for d in depts]})

    @http.route("/api/agents", type="http", auth="none", csrf=False, methods=["GET", "OPTIONS"])
    def agents(self, **kwargs):
        if request.httprequest.method == "OPTIONS": return _cors_preflight_response()
        env, err = _require_auth()
        if err: return err
        assistants = _model(env, "llm.assistant")
        data = [{"id": "role-default", "name": _assigned_agent(env, env.user), "description": "Agent اختصاصی نقش شما؛ محدود به Capabilityهای مؤثر همان کاربر.", "tools": 0, "assigned_to_user": True}]
        if assistants is not None:
            for a in assistants.search([], order="name"):
                data.append({"id": a.id, "name": a.name, "description": getattr(a, "description", "") or "", "tools": len(a.tool_ids) if hasattr(a, "tool_ids") else 0})
        return _json_response({"agents": data})

    @http.route("/api/tasks", type="http", auth="none", csrf=False, methods=["GET", "POST", "OPTIONS"])
    def tasks(self, **kwargs):
        if request.httprequest.method == "OPTIONS": return _cors_preflight_response()
        env, err = _require_auth()
        if err: return err
        model = _model(env, "project.task")
        if model is None: return _json_response({"tasks": []})
        if request.httprequest.method == "POST":
            try:
                payload = json.loads(request.httprequest.data or b"{}")
            except (TypeError, ValueError):
                return _json_response({"error": "invalid JSON body"}, status=400)
            if not isinstance(payload, dict):
                return _json_response({"error": "JSON body must be an object"}, status=400)
            vals = {"name": payload.get("name"), "description": payload.get("description", "")}
            if not vals["name"]: return _json_response({"error": "name is required"}, status=400)
            try:
                result = env["ai.gateway.execution.gate"].execute(
                    "project.task.create", args={"name": str(vals["name"])[:200], "description": str(vals["description"])[:4000]}
                )
            except Exception:  # noqa: BLE001
                _logger.exception("task creation failed")
                return _json_response({"error": "task could not be created"}, status=409)
            task = model.browse(result.get("record_id")).exists()
            _audit(env, env.user.id, "experience_api", "task.create", {"task_id": task.id if task else None}, True)
            return _json_response({"id": task.id if task else result.get("record_id"), "name": task.name if task else vals["name"]}, status=201)
        tasks = model.search([("create_uid", "=", env.user.id)], order="create_date desc", limit=100)
        return _json_response({"tasks": [{"id": t.id, "name": t.name, "description": t.description or "", "state": getattr(t.stage_id, "name", "") if hasattr(t, "stage_id") else ""} for t in tasks]})

    @http.route("/api/calendar", type="http", auth="none", csrf=False, methods=["GET", "POST", "OPTIONS"])
    def calendar(self, **kwargs):
        if request.httprequest.method == "OPTIONS": return _cors_preflight_response()
        env, err = _require_auth()
        if err: return err
        model = _model(env, "calendar.event")
        if model is None: return _json_response({"events": []})
        if request.httprequest.method == "POST":
            try:
                payload = json.loads(request.httprequest.data or b"{}")
            except (TypeError, ValueError):
                return _json_response({"error": "invalid JSON body"}, status=400)
            if not isinstance(payload, dict) or not payload.get("name") or not payload.get("start"):
                return _json_response({"error": "name and start are required"}, status=400)
            try:
                result = env["ai.gateway.execution.gate"].execute(
                    "calendar.event.create", args={
                        "name": str(payload["name"])[:200],
                        "start": str(payload["start"]),
                        "stop": str(payload.get("stop") or payload["start"]),
                        "allday": bool(payload.get("allday")),
                        "description": str(payload.get("description") or "")[:4000],
                        "location": str(payload.get("location") or "")[:500],
                    },
                )
                _audit(env, env.user.id, "experience_api", "calendar.create", {"record_id": result.get("record_id")}, True)
                return _json_response(result, status=201)
            except Exception:  # noqa: BLE001
                _logger.exception("calendar event creation failed")
                return _json_response({"error": "calendar event could not be created"}, status=409)
        start = kwargs.get("start")
        end = kwargs.get("end")
        domain = ["|", ("partner_ids", "in", [env.user.partner_id.id]), ("user_id", "=", env.user.partner_id.id)]
        if start: domain.append(("stop", ">=", start))
        if end: domain.append(("start", "<=", end))
        events = model.search(domain, order="start asc", limit=250)
        return _json_response({"events": [{
            "id": event.id, "name": event.name, "start": str(event.start) if event.start else None,
            "stop": str(event.stop) if event.stop else None, "allday": bool(event.allday),
            "location": event.location or "", "description": event.description or "",
        } for event in events]})

    @http.route("/api/notifications", type="http", auth="none", csrf=False, methods=["GET", "OPTIONS"])
    def notifications(self, **kwargs):
        if request.httprequest.method == "OPTIONS": return _cors_preflight_response()
        env, err = _require_auth()
        if err: return err
        model = _model(env, "mail.notification")
        out=[]
        if model is not None:
            recs = model.search([("res_partner_id", "=", env.user.partner_id.id)], order="id desc", limit=100)
            for r in recs:
                msg = r.mail_message_id
                out.append({"id": r.id, "subject": msg.subject or "", "body": msg.body or "", "is_read": getattr(r, "is_read", False), "date": str(msg.date or "")})
        return _json_response({"notifications": out})

    @http.route("/api/models", type="http", auth="none", csrf=False, methods=["GET", "OPTIONS"])
    def models(self, **kwargs):
        if request.httprequest.method == "OPTIONS": return _cors_preflight_response()
        env, err = _require_auth()
        if err: return err
        model = _model(env, "ai.model.profile")
        out=[]
        if model is not None:
            labels = {
                "chat": "دستیار گفتگو", "reasoning": "دستیار تحلیل",
                "vision": "تحلیل تصویر", "embedding": "جستجوی دانش",
            }
            for m in model.sudo().search([], order="purpose,name"):
                # Keep infrastructure/provider/model identifiers out of the
                # customer-facing capability surface.
                out.append({"label": labels.get(m.purpose, "قابلیت هوشمند"),
                            "purpose": m.purpose, "available": bool(m.active),
                            "production": bool(m.production)})
        return _json_response({"models": out})


    @http.route("/api/files/analyze", type="http", auth="none", csrf=False, methods=["POST", "OPTIONS"])
    def analyze_file(self, **kwargs):
        if request.httprequest.method == "OPTIONS": return _cors_preflight_response()
        env, err = _require_auth()
        if err: return err
        try:
            payload = json.loads(request.httprequest.data or b"{}")
            if not isinstance(payload, dict):
                return _json_response({"error": "JSON body must be an object"}, status=400)
            filename = _safe_name(payload.get("filename", "file"), "file")
            raw = base64.b64decode(payload.get("data_base64", ""), validate=True)
            validate_upload(filename, raw, max_bytes=MAX_UPLOAD_BYTES)
            question = str(payload.get("question") or "خلاصه و نکات مهم این فایل را توضیح بده")[:4000]
            lower = filename.lower()
            text = ""
            if lower.endswith((".txt", ".md", ".csv", ".json")):
                text = raw.decode("utf-8", errors="replace")
            elif lower.endswith(".pdf"):
                try:
                    from pypdf import PdfReader
                    reader = PdfReader(io.BytesIO(raw))
                    if len(reader.pages) > 100:
                        return _json_response({"error": "PDF exceeds the page limit"}, status=413)
                    chunks = []
                    for page in reader.pages:
                        chunks.append(page.extract_text() or "")
                        if sum(len(chunk) for chunk in chunks) >= 60000:
                            break
                    text = "\n".join(chunks)
                except ImportError:
                    return _json_response({"error": "PDF reader is not installed"}, status=501)
            elif lower.endswith(".docx"):
                try:
                    from docx import Document
                    doc = Document(io.BytesIO(raw)); text = "\n".join(p.text for p in doc.paragraphs)
                except ImportError:
                    return _json_response({"error": "DOCX reader is not installed"}, status=501)
            elif lower.endswith(".xls"):
                return _json_response({"error": "legacy XLS extraction is not available"}, status=501)
            elif lower.endswith(".xlsx"):
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                    chunks=[]
                    for ws in wb.worksheets[:5]:
                        chunks.append(f"[Sheet: {ws.title}]")
                        for row in ws.iter_rows(max_row=100, values_only=True): chunks.append(" | ".join("" if v is None else str(v) for v in row))
                    text = "\n".join(chunks)
                except ImportError:
                    return _json_response({"error": "XLSX reader is not installed"}, status=501)
            elif lower.endswith((".png", ".jpg", ".jpeg", ".webp")):
                import requests
                profile = env["ai.model.router"].route(purpose="vision")
                base = profile.endpoint or env["ir.config_parameter"].sudo().get_param("company_ai_demo.vision_api_base", "http://127.0.0.1:8001/v1")
                mime = "image/png" if lower.endswith(".png") else "image/jpeg"
                resp = requests.post(f"{base.rstrip('/')}/chat/completions", json={"model":profile.model_id,"messages":[{"role":"user","content":[{"type":"text","text":question},{"type":"image_url","image_url":{"url":f"data:{mime};base64,{base64.b64encode(raw).decode()}"}}]}],"max_tokens":1500}, timeout=90)
                resp.raise_for_status(); answer=resp.json()["choices"][0]["message"]["content"]
                return _json_response({"filename": filename, "kind": "vision", "analysis": scrub_public_text(answer)})
            else:
                return _json_response({"error": "unsupported file type"}, status=415)
            text = text[:60000]
            return _json_response({"filename": filename, "kind": "text", "analysis": scrub_public_text(text if text else "No extractable text was found."), "question": scrub_public_text(question)})
        except (TypeError, ValueError):
            return _json_response({"error": "invalid or unsupported file upload"}, status=400)
        except Exception:
            _logger.exception("file analysis failed")
            return _json_response({"error": "file analysis failed, try again later"}, status=500)


    @http.route("/api/artifacts/<int:attachment_id>", type="http", auth="none", csrf=False, methods=["GET", "OPTIONS"])
    def download_artifact(self, attachment_id, **kwargs):
        if request.httprequest.method == "OPTIONS": return _cors_preflight_response()
        env, err = _require_auth()
        if err: return err
        att = env["ir.attachment"].sudo().browse(attachment_id)
        if not att.exists() or att.res_model != "res.users" or att.res_id != env.user.id:
            return _json_response({"error": "artifact not found or access denied"}, status=404)
        data = base64.b64decode(att.datas or b"")
        headers = [
            ("Content-Type", att.mimetype or "application/octet-stream"),
            ("Content-Disposition", "attachment; filename*=UTF-8''%s" % _quote_filename(att.name or "artifact")),
            ("Content-Length", str(len(data))),
        ]
        return request.make_response(data, headers=headers)

    @http.route("/api/artifacts/generate", type="http", auth="none", csrf=False, methods=["POST", "OPTIONS"])
    def generate_artifact(self, **kwargs):
        if request.httprequest.method == "OPTIONS": return _cors_preflight_response()
        env, err = _require_auth()
        if err: return err
        try:
            payload = json.loads(request.httprequest.data or b"{}")
            kind = payload.get("type", "csv")
            filename, mimetype, data = _artifact_bytes(kind, payload.get("title", "artifact"), payload)
            _audit(env, env.user.id, "experience_api", "artifact.generate", {"type": kind, "filename": filename}, True)
            return _json_response({"filename": filename, "mimetype": mimetype, "data_base64": base64.b64encode(data).decode("ascii")})
        except (ValueError, RuntimeError, ImportError) as exc:
            _audit(env, env.user.id, "experience_api", "artifact.generate", {"type": payload.get("type") if isinstance(payload, dict) else ""}, False, str(exc))
            detail = "unsupported artifact type or missing data" if isinstance(exc, (ValueError, ImportError)) else "artifact generation failed, try again later"
            return _json_response({"error": detail}, status=400 if isinstance(exc, ValueError) else 501)
