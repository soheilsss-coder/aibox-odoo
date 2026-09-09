"""Local, bounded document extraction shared by chat attachments and RAG.

This module deliberately has no Odoo imports.  It can therefore be tested in
isolation and can be reused by the native document worker later.  Heavy parser
and OCR dependencies are imported lazily: a partial installation fails with a
specific parser error instead of breaking the Odoo registry at import time.
"""
from __future__ import annotations

import csv
import html
import io
import json
import logging
import mimetypes
import os
import re
import tempfile
import warnings
from importlib import metadata as importlib_metadata
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from threading import Lock

_logger = logging.getLogger(__name__)

# PyMuPDF/RapidOCR may import SWIG extension types that still emit
# Python 3.12 DeprecationWarning messages during Odoo registry loading.
# They are third-party compatibility noise, not actionable product warnings.
warnings.filterwarnings(
    "ignore",
    message=r"builtin type (SwigPyPacked|SwigPyObject|swigvarlink) has no __module__ attribute",
    category=DeprecationWarning,
)

# Unstructured's optional telemetry is enabled by default in recent releases.
# The appliance is explicitly local/on-prem, so disable analytics before any
# lazy parser import can initialize that package.
os.environ.setdefault("DO_NOT_TRACK", "1")
os.environ.setdefault("SCARF_NO_ANALYTICS", "1")

SUPPORTED_EXTENSIONS = frozenset({
    ".txt", ".md", ".csv", ".json", ".html", ".htm",
    ".pdf", ".docx", ".xlsx", ".pptx",
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff",
})

_DOCLING_EXTENSIONS = frozenset({
    ".pdf", ".docx", ".xlsx", ".pptx", ".html", ".htm",
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff",
})

_MAX_BLOCKS = 100_000
_DOCling_LOCK = Lock()
_DOCling_CONVERTER = None


class ExtractionError(RuntimeError):
    """A safe, user-facing extraction failure with a stable error code."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class ExtractionResult:
    text: str
    blocks: list[dict] = field(default_factory=list)
    parser: str = "text"
    parser_version: str = "builtin"
    content_type: str = "text/plain"
    page_count: int = 0
    warnings: list[str] = field(default_factory=list)
    checksum: str = ""

    @property
    def char_count(self):
        return len(self.text)


def _bounded_text(text, max_chars):
    text = re.sub(r"\n{3,}", "\n\n", text or "").strip()
    if len(text) > max_chars:
        raise ExtractionError("extracted_text_limit", "extracted document text exceeds the safe processing limit")
    return text


def _block(text, **metadata):
    text = re.sub(r"\n{3,}", "\n\n", str(text or "")).strip()
    if not text:
        return None
    value = {"text": text}
    value.update({k: v for k, v in metadata.items() if v not in (None, "")})
    return value


def _result_from_blocks(blocks, **kwargs):
    blocks = [b for b in blocks if b and b.get("text")][: _MAX_BLOCKS]
    text = "\n\n".join(b["text"] for b in blocks)
    return ExtractionResult(text=text, blocks=blocks, **kwargs)


class _SafeHTMLParser(HTMLParser):
    """Extract visible HTML text without indexing scripts or styles."""

    _ignored = {"script", "style", "noscript", "template", "svg"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self._depth = 0

    def handle_starttag(self, tag, attrs):
        if tag.lower() in self._ignored:
            self._depth += 1
        elif tag.lower() in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag.lower() in self._ignored and self._depth:
            self._depth -= 1
        elif tag.lower() in {"p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._depth:
            self.parts.append(data)


def _read_utf8(path):
    try:
        return Path(path).read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError("invalid_encoding", "text document must be UTF-8") from exc


def _extract_plain(path, ext, max_chars):
    if ext in {".txt", ".md"}:
        text = _read_utf8(path)
        return _result_from_blocks(
            [_block(text, content_type="markdown" if ext == ".md" else "text")],
            parser="builtin-text", parser_version="1", content_type="text/markdown" if ext == ".md" else "text/plain",
        )
    if ext in {".html", ".htm"}:
        parser = _SafeHTMLParser()
        try:
            parser.feed(_read_utf8(path))
            parser.close()
        except (UnicodeDecodeError, ValueError) as exc:
            raise ExtractionError("invalid_html", "HTML document could not be decoded") from exc
        text = html.unescape(" ".join(" ".join(parser.parts).split()))
        return _result_from_blocks(
            [_block(text, content_type="html")],
            parser="builtin-html", parser_version="1", content_type="text/html",
        )
    if ext == ".json":
        try:
            data = json.loads(_read_utf8(path))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ExtractionError("invalid_json", "JSON document is not valid UTF-8 JSON") from exc
        lines = []

        def visit(value, prefix="$", depth=0):
            if depth > 40:
                lines.append(f"{prefix}: [nested value omitted]")
                return
            if isinstance(value, dict):
                for key, item in value.items():
                    visit(item, f"{prefix}.{key}", depth + 1)
            elif isinstance(value, list):
                for index, item in enumerate(value):
                    visit(item, f"{prefix}[{index}]", depth + 1)
            else:
                lines.append(f"{prefix}: {value}")

        visit(data)
        text = "\n".join(lines) or "$: (empty JSON value)"
        return _result_from_blocks(
            [_block(text, content_type="json")],
            parser="builtin-json", parser_version="1", content_type="application/json",
        )
    if ext == ".csv":
        raw = _read_utf8(path)
        try:
            rows = csv.reader(io.StringIO(raw, newline=""))
            lines = []
            for row_number, row in enumerate(rows, start=1):
                if row_number > 100_000:
                    raise ExtractionError("row_limit", "CSV row count exceeds the safe processing limit")
                lines.append(f"row {row_number}: " + " | ".join(cell.strip() for cell in row))
        except csv.Error as exc:
            raise ExtractionError("invalid_csv", "CSV document is malformed") from exc
        return _result_from_blocks(
            [_block("\n".join(lines), content_type="csv")],
            parser="builtin-csv", parser_version="1", content_type="text/csv",
        )
    raise ExtractionError("unsupported_type", f"unsupported extraction type: {ext}")


def get_ocr():
    from rapidocr_onnxruntime import RapidOCR
    return RapidOCR()


def ocr_image_file(image_path):
    try:
        result, _ = get_ocr()(image_path)
    except ImportError as exc:
        raise ExtractionError("ocr_dependency_missing", "OCR engine is not installed") from exc
    except Exception as exc:  # noqa: BLE001
        _logger.warning("OCR failed on %s: %s", image_path, exc)
        return ""
    if not result:
        return ""
    return "\n".join(str(line[1]) for line in result if len(line) > 1)


def ocr_scanned_pdf(pdf_path, max_pages=500):
    try:
        import fitz
    except ImportError as exc:
        raise ExtractionError("pdf_dependency_missing", "PDF/OCR dependencies are not installed") from exc
    texts = []
    try:
        with fitz.open(pdf_path) as doc:
            for page_number, page in enumerate(doc, start=1):
                if page_number > max_pages:
                    raise ExtractionError("page_limit", "PDF page count exceeds the safe processing limit")
                pix = page.get_pixmap(dpi=200, alpha=False)
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                    temp_name = tmp.name
                try:
                    pix.save(temp_name)
                    text = ocr_image_file(temp_name)
                    if text:
                        texts.append(_block(text, page=page_number, content_type="ocr"))
                finally:
                    try:
                        os.unlink(temp_name)
                    except FileNotFoundError:
                        pass
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("pdf_extract_failed", "PDF OCR failed") from exc
    return _result_from_blocks(
        texts,
        parser="rapidocr-pdf", parser_version="1", content_type="application/pdf",
        page_count=len(texts),
    )


def _docling_converter():
    global _DOCling_CONVERTER
    if _DOCling_CONVERTER is not None:
        return _DOCling_CONVERTER
    with _DOCling_LOCK:
        if _DOCling_CONVERTER is None:
            try:
                from docling.document_converter import DocumentConverter
            except ImportError as exc:
                raise ExtractionError("docling_dependency_missing", "Docling is not installed") from exc
            _DOCling_CONVERTER = DocumentConverter()
    return _DOCling_CONVERTER


def _docling_extract(path, ext, max_chars):
    converter = _docling_converter()
    try:
        conversion = converter.convert(path)
        document = conversion.document
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("docling_extract_failed", "Docling could not parse the document") from exc

    blocks = []
    # Docling versions expose text items differently. Keep the introspection
    # defensive so a minor upstream shape change degrades to Markdown rather
    # than breaking ingestion entirely.
    for item in list(getattr(document, "texts", []) or []):
        text = getattr(item, "text", "")
        provenance = list(getattr(item, "prov", []) or [])
        prov = provenance[0] if provenance else None
        page = getattr(prov, "page_no", None) if prov else None
        bbox = getattr(prov, "bbox", None) if prov else None
        label = getattr(item, "label", None)
        # Keep parser-native labels and coordinates where available.  The
        # canonical chunker stores these as citation metadata without making
        # assumptions about one Docling release's concrete geometry class.
        blocks.append(_block(
            text,
            page=page,
            content_type=str(label or "text"),
            coordinates=str(bbox) if bbox else None,
            section=getattr(item, "section", None) or getattr(item, "heading", None),
            table=getattr(item, "table_name", None) or getattr(item, "table_id", None),
            sheet=(getattr(item, "sheet_name", None) or (getattr(prov, "sheet_name", None) if prov else None)),
            slide=(getattr(item, "slide_no", None) or (getattr(prov, "slide_no", None) if prov else None)),
        ))
    if not blocks:
        try:
            markdown = document.export_to_markdown()
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError("docling_export_failed", "Docling returned no usable text") from exc
        blocks = [_block(markdown, content_type="document")]
    try:
        parser_version = importlib_metadata.version("docling-core")
    except importlib_metadata.PackageNotFoundError:
        parser_version = str(getattr(conversion, "version", "unknown"))
    result = _result_from_blocks(
        blocks,
        parser="docling", parser_version=parser_version,
        content_type=mimetypes.guess_type(str(path))[0] or "application/octet-stream",
        page_count=len(getattr(document, "pages", {}) or {}),
    )
    # A document with image-only pages can have no text items. Let the OCR
    # fallback handle it instead of activating an empty revision.
    if not result.text.strip() and ext == ".pdf":
        return ocr_scanned_pdf(path)
    return result


def _unstructured_extract(path, ext, max_chars):
    # Unstructured 0.27 may self-download en_core_web_sm while classifying
    # office paragraphs. That is not acceptable on an offline appliance. Use
    # its fallback only when the pinned model is already installed locally;
    # the deterministic adapters below handle the same formats otherwise.
    if ext in {".docx", ".xlsx", ".pptx"}:
        try:
            import spacy
            if not spacy.util.is_package("en_core_web_sm"):
                raise ExtractionError(
                    "unstructured_model_missing",
                    "Unstructured local sentence model is not installed",
                )
        except ImportError as exc:
            raise ExtractionError("unstructured_dependency_missing", "Unstructured is not installed") from exc
    try:
        from unstructured.partition.auto import partition
    except ImportError as exc:
        raise ExtractionError("unstructured_dependency_missing", "Unstructured is not installed") from exc
    try:
        elements = partition(filename=str(path), strategy="auto")
    except TypeError:
        # Older Unstructured versions do not accept strategy on every backend.
        try:
            elements = partition(filename=str(path))
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError("unstructured_extract_failed", "Unstructured could not parse the document") from exc
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("unstructured_extract_failed", "Unstructured could not parse the document") from exc

    blocks = []
    pages = set()
    for element in elements[:_MAX_BLOCKS]:
        metadata = getattr(element, "metadata", None)
        page = getattr(metadata, "page_number", None) if metadata else None
        if page is not None:
            pages.add(page)
        category = getattr(element, "category", None)
        coordinates = getattr(metadata, "coordinates", None) if metadata else None
        section = getattr(metadata, "section", None) if metadata else None
        table = getattr(metadata, "table_name", None) if metadata else None
        sheet = (
            getattr(metadata, "sheet_name", None) or getattr(metadata, "page_name", None)
            if metadata else None
        )
        slide = getattr(metadata, "slide_number", None) if metadata else None
        block = _block(
            str(element), page=page, section=str(section) if section else None,
            content_type=category,
            coordinates=str(coordinates) if coordinates else None,
            table=str(table) if table else None,
            sheet=str(sheet) if sheet else None,
            slide=slide,
        )
        if block:
            blocks.append(block)
    try:
        parser_version = importlib_metadata.version("unstructured")
    except importlib_metadata.PackageNotFoundError:
        parser_version = "unknown"
    result = _result_from_blocks(
        blocks,
        parser="unstructured", parser_version=parser_version,
        content_type=mimetypes.guess_type(str(path))[0] or "application/octet-stream",
        page_count=len(pages),
    )
    if ext == ".pdf" and len(result.text.strip()) < 30:
        return ocr_scanned_pdf(path)
    if ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"} and not result.text.strip():
        text = ocr_image_file(path)
        return _result_from_blocks(
            [_block(text, content_type="ocr")],
            parser="rapidocr", parser_version="1", content_type=result.content_type,
        )
    return result


def _local_structured_extract(path, ext, max_chars):
    """Offline fallback for parser installations without Unstructured's model.

    Docling remains the primary structured parser and Unstructured remains the
    secondary parser when its local model is present. These small adapters are
    not an AI implementation: they prevent a missing optional NLP model from
    turning a native installation into a network-dependent ingestion path.
    """
    if ext == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ExtractionError("office_dependency_missing", "DOCX parser dependencies are not installed") from exc
        document = Document(path)
        blocks = []
        for paragraph in document.paragraphs:
            style = getattr(getattr(paragraph, "style", None), "name", "")
            block = _block(
                paragraph.text,
                content_type="title" if style.lower().startswith("heading") else "paragraph",
                section=style if style.lower().startswith("heading") else None,
            )
            if block:
                blocks.append(block)
        for table_number, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows[:10_000]:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            block = _block(
                "\n".join(rows),
                content_type="table",
                table=f"table-{table_number}",
            )
            if block:
                blocks.append(block)
        return _result_from_blocks(
            blocks,
            parser="python-docx", parser_version=importlib_metadata.version("python-docx"),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ExtractionError("office_dependency_missing", "XLSX parser dependencies are not installed") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        blocks = []
        try:
            for sheet in workbook.worksheets:
                rows = []
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if row_number > 100_000:
                        raise ExtractionError("row_limit", "spreadsheet row count exceeds the safe processing limit")
                    values = [str(value).strip() for value in row if value is not None]
                    if values:
                        rows.append(" | ".join(values))
                block = _block("\n".join(rows), content_type="sheet", sheet=sheet.title)
                if block:
                    blocks.append(block)
        finally:
            workbook.close()
        return _result_from_blocks(
            blocks,
            parser="openpyxl", parser_version=importlib_metadata.version("openpyxl"),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    if ext == ".pptx":
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ExtractionError("office_dependency_missing", "PPTX parser dependencies are not installed") from exc
        presentation = Presentation(path)
        blocks = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    parts.append(shape.text)
                if getattr(shape, "has_table", False):
                    parts.extend(
                        " | ".join(cell.text.strip() for cell in row.cells)
                        for row in shape.table.rows[:10_000]
                    )
            block = _block("\n".join(parts), content_type="slide", slide=slide_number)
            if block:
                blocks.append(block)
        return _result_from_blocks(
            blocks,
            parser="python-pptx", parser_version=importlib_metadata.version("python-pptx"),
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

    if ext == ".pdf":
        try:
            import fitz
        except ImportError as exc:
            raise ExtractionError("pdf_dependency_missing", "PDF parser dependencies are not installed") from exc
        blocks = []
        try:
            with fitz.open(path) as document:
                for page_number, page in enumerate(document, start=1):
                    text = page.get_text("text")
                    block = _block(text, page=page_number, content_type="pdf-text")
                    if block:
                        blocks.append(block)
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError("pdf_extract_failed", "PDF text extraction failed") from exc
        if blocks:
            return _result_from_blocks(
                blocks,
                parser="pymupdf", parser_version=importlib_metadata.version("PyMuPDF"),
                content_type="application/pdf", page_count=len(blocks),
            )
        return ocr_scanned_pdf(path)

    raise ExtractionError("unsupported_type", f"no offline fallback for {ext}")


def extract_file(path, filename=None, max_chars=2_000_000, language_hints=None):
    """Extract a bounded, local document into canonical blocks.

    `language_hints` is reserved for the parser worker; it is intentionally
    not silently passed to every optional backend because Unstructured and
    Docling changed parameter names across releases.
    """
    del language_hints
    path = Path(path)
    name = filename or path.name
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError("unsupported_type", f"unsupported file type: {ext or 'unknown'}")
    if not path.is_file():
        raise ExtractionError("missing_file", "document payload is missing")

    if ext in {".txt", ".md", ".csv", ".json", ".html", ".htm"}:
        result = _extract_plain(path, ext, max_chars)
    elif ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        try:
            text = ocr_image_file(path)
        except ExtractionError:
            # If Docling is installed it may have a better image backend.
            result = _docling_extract(path, ext, max_chars)
        else:
            if text.strip():
                result = _result_from_blocks(
                    [_block(text, content_type="ocr")],
                    parser="rapidocr", parser_version="1",
                    content_type=mimetypes.guess_type(str(path))[0] or "image/*",
                )
            else:
                try:
                    result = _docling_extract(path, ext, max_chars)
                except ExtractionError:
                    result = _result_from_blocks(
                        [], parser="rapidocr", parser_version="1",
                        content_type=mimetypes.guess_type(str(path))[0] or "image/*",
                    )
    else:
        try:
            result = _docling_extract(path, ext, max_chars)
        except ExtractionError as docling_error:
            try:
                result = _unstructured_extract(path, ext, max_chars)
            except ExtractionError as fallback_error:
                try:
                    result = _local_structured_extract(path, ext, max_chars)
                except ExtractionError as local_error:
                    # Preserve the most actionable dependency/parser error,
                    # while retaining the parser chain in the exception cause.
                    raise local_error from fallback_error

    result.text = _bounded_text(result.text, max_chars)
    if not result.text:
        raise ExtractionError("ocr_empty", "document contains no extractable text")
    return result
